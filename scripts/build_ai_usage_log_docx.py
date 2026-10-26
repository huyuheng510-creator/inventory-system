#!/usr/bin/env python3
"""Render docs/AI_USAGE_LOG.md as a report-ready Word appendix.

The Markdown file is the single source of truth: this converter turns its
headings, paragraphs and pipe tables into a formatted A4 DOCX so the text in
the report and the Markdown cannot drift apart.
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

PROJECT = Path(__file__).resolve().parents[1]
SOURCE = PROJECT / "docs" / "AI_USAGE_LOG.md"
TARGET = PROJECT / "docs" / "AI_Usage_Log.docx"

HEADING_COLOUR = RGBColor(0x1F, 0x38, 0x64)
BODY_FONT = "Times New Roman"


def shade(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_field(paragraph, instruction: str):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def add_runs(paragraph, text: str, size: float | None = None, bold_all: bool = False):
    """Add text to a paragraph, honouring **bold** segments."""
    text = text.replace("`", "").replace("✓", "Yes")
    for index, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if not part:
            continue
        run = paragraph.add_run(part)
        run.font.name = BODY_FONT
        if size:
            run.font.size = Pt(size)
        run.bold = bold_all or index % 2 == 1


def add_table(document: Document, rows: list[list[str]]):
    header, body = rows[0], rows[1:]
    columns = len(header)
    table = document.add_table(rows=1, cols=columns)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    widths = {
        3: [Inches(1.55), Inches(1.9), Inches(2.75)],
        4: [Inches(1.35), Inches(1.0), Inches(1.5), Inches(2.35)],
        5: [Inches(0.32), Inches(0.82), Inches(1.45), Inches(1.65), Inches(1.95)],
    }.get(columns, [Inches(6.2 / columns)] * columns)

    header_cells = table.rows[0].cells
    for i, value in enumerate(header):
        header_cells[i].width = widths[i]
        paragraph = header_cells[i].paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        add_runs(paragraph, value.strip(), size=8.5, bold_all=True)
        shade(header_cells[i], "D9E2F3")

    for row in body:
        cells = table.add_row().cells
        for i in range(columns):
            value = row[i] if i < len(row) else ""
            cells[i].width = widths[i]
            paragraph = cells[i].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            add_runs(paragraph, value.strip(), size=8.5)

    # A fixed layout with explicit grid columns keeps the printed widths stable
    # in Word and LibreOffice, and no row may be split over two pages.
    table_pr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    table_pr.append(layout)
    for grid_col, width in zip(table._tbl.tblGrid.gridCol_lst, widths):
        grid_col.set(qn("w:w"), str(int(width.inches * 1440)))
    for index, row in enumerate(table.rows):
        tr_pr = row._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        if index == 0:
            tr_pr.append(OxmlElement("w:tblHeader"))
        for i, cell in enumerate(row.cells):
            cell.width = widths[i]
    document.add_paragraph()


def parse_markdown(text: str):
    """Yield ('heading', level, text) / ('para', text) / ('table', rows)."""
    blocks = []
    paragraph_lines: list[str] = []
    table_lines: list[str] = []

    def flush_paragraph():
        if paragraph_lines:
            blocks.append(("para", " ".join(paragraph_lines)))
            paragraph_lines.clear()

    def flush_table():
        if table_lines:
            rows = []
            for line in table_lines:
                cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
                if all(re.fullmatch(r":?-{2,}:?", cell) for cell in cells if cell):
                    continue
                rows.append(cells)
            if rows:
                blocks.append(("table", rows))
            table_lines.clear()

    for raw in text.splitlines():
        line = raw.strip()
        if line.startswith("|"):
            flush_paragraph()
            table_lines.append(line)
            continue
        flush_table()
        if line.startswith("#"):
            flush_paragraph()
            level = len(line) - len(line.lstrip("#"))
            blocks.append(("heading", level, line[level:].strip()))
        elif not line:
            flush_paragraph()
        else:
            paragraph_lines.append(line)
    flush_table()
    flush_paragraph()
    return blocks


def main():
    blocks = parse_markdown(SOURCE.read_text(encoding="utf-8"))
    document = Document()

    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.left_margin = Inches(0.98)
    section.right_margin = Inches(0.98)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)

    normal = document.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_runs(header, "COMP603/ENSE600 Project 2 — StockFlow AI Usage Log", size=8.5)
    header.runs[0].font.color.rgb = RGBColor(0x70, 0x70, 0x70)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(footer, "PAGE")

    title_added = False
    for block in blocks:
        if block[0] == "heading":
            _, level, text = block
            if not title_added and level == 1:
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                add_runs(paragraph, text, size=16, bold_all=True)
                for run in paragraph.runs:
                    run.font.color.rgb = HEADING_COLOUR
                paragraph.paragraph_format.space_after = Pt(10)
                title_added = True
            else:
                paragraph = document.add_paragraph()
                add_runs(paragraph, text, size=13 if level == 2 else 12, bold_all=True)
                for run in paragraph.runs:
                    run.font.color.rgb = HEADING_COLOUR
                paragraph.paragraph_format.space_before = Pt(12)
                paragraph.paragraph_format.space_after = Pt(5)
        elif block[0] == "para":
            paragraph = document.add_paragraph()
            add_runs(paragraph, block[1])
        else:
            add_table(document, block[1])

    properties = document.core_properties
    properties.title = "Appendix A: AI Usage Log (Project 2)"
    properties.author = "Yuheng Hu (26366447)"
    properties.subject = "COMP603/ENSE600 Project 2 - StockFlow Inventory Management System"

    document.save(TARGET)

    reopened = Document(TARGET)
    print("saved", TARGET)
    print("paragraphs:", len(reopened.paragraphs), "tables:", len(reopened.tables))
    print("table rows:", [len(table.rows) for table in reopened.tables])
    print("first paragraph:", reopened.paragraphs[0].text[:100])


if __name__ == "__main__":
    main()
