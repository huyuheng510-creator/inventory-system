#!/usr/bin/env python3
"""Build a report-ready Word draft of the four Project 2 design figures.

The images in docs/design are inserted at their report width (6.2 in) with the
captions and short placement notes so the figures can be copied into the Word
report without re-doing the layout.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

PROJECT = Path(__file__).resolve().parents[1]
DESIGN = PROJECT / "docs" / "design"
TARGET = PROJECT / "docs" / "P2_Design_Figures.docx"

BODY_FONT = "Times New Roman"
HEADING_COLOUR = RGBColor(0x1F, 0x38, 0x64)
CAPTION_COLOUR = RGBColor(0x40, 0x40, 0x40)


def body(document: Document, text: str, size: float = 11.0):
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.font.name = BODY_FONT
    run.font.size = Pt(size)
    return paragraph


def heading(document: Document, text: str, level: int = 1, page_break: bool = False):
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.font.name = BODY_FONT
    run.font.size = Pt(15 if level == 1 else 12)
    run.bold = True
    run.font.color.rgb = HEADING_COLOUR
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.page_break_before = page_break
    return paragraph


def caption(document: Document, text: str):
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.font.name = BODY_FONT
    run.font.size = Pt(10)
    run.italic = True
    run.font.color.rgb = CAPTION_COLOUR
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(10)
    return paragraph


def main():
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.left_margin = Inches(0.98)
    section.right_margin = Inches(0.98)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)

    normal = document.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    title = document.add_paragraph()
    run = title.add_run("Project 2 Design Figures (Report-Ready Draft)")
    run.font.name = BODY_FONT
    run.font.size = Pt(18)
    run.bold = True
    run.font.color.rgb = HEADING_COLOUR
    title.paragraph_format.space_after = Pt(4)

    subtitle = document.add_paragraph()
    run = subtitle.add_run("StockFlow Inventory Management System - "
                           "COMP603/ENSE600 Project 2")
    run.font.name = BODY_FONT
    run.font.size = Pt(11.5)
    run.italic = True

    body(document,
         "The document contains four figures: Figure 1 (use-case diagram), Figure 2 (final "
         "class diagram), Figure 3 (entity-relationship diagram of the Derby schema) and "
         "Figure 4 (MVC and layered architecture). The PNG files and the reproducible Pillow "
         "generator are in docs/design/. The class diagram and the ER diagram are checked "
         "against the built code and src/main/resources/db/schema.sql; the use-case model is "
         "checked against the GUI tabs, controllers and service methods.")

    contents = document.add_paragraph()
    contents.paragraph_format.space_before = Pt(14)
    run = contents.add_run("Figures in this document")
    run.bold = True
    run.font.name = BODY_FONT
    run.font.size = Pt(11)
    for item in [
        "Figure 1 - Use-case diagram (Project 2 scope, UC-01 to UC-18)",
        "Figure 2 - Final class diagram (ui, service, dao, model)",
        "Figure 3 - Entity-relationship diagram of the embedded Derby database",
        "Figure 4 - MVC and layered architecture with the design patterns",
    ]:
        paragraph = body(document, item, size=10.5)
        paragraph.paragraph_format.space_after = Pt(2)

    heading(document, "Figure 1 - Use case diagram", page_break=True)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(DESIGN / "figure1_use_case_p2.png"),
                                    width=Inches(6.2))
    caption(document, "Figure 1: Use-case diagram for the Project 2 StockFlow system "
                      "(UC-15 to UC-18 are the Project 2 additions).")
    body(document,
         "Recommended placement: Section 3 (Software Design), use-case subsection; the "
         "model extends the Project 1 use cases with UC-15 to UC-18. See "
         "docs/design/README.md for the full relationship notes.", size=10.5)

    heading(document, "Figure 2 - Final class diagram", page_break=True)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(DESIGN / "figure2_class_diagram_p2.png"),
                                    width=Inches(6.2))
    caption(document, "Figure 2: Final class diagram of the Project 2 design (MVC layer, "
                      "service layer, JDBC DAO layer and domain model).")
    body(document,
         "Recommended placement: Section 3, class/design subsection (replaces the Project 1 "
         "iteration-2 class diagram). See docs/design/README.md for the suggested "
         "design-evolution paragraph.", size=10.5)

    heading(document, "Figure 3 - Entity-relationship diagram", page_break=True)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(DESIGN / "figure3_er_diagram_p2.png"),
                                    width=Inches(6.2))
    caption(document, "Figure 3: Entity-relationship diagram of the embedded Apache Derby "
                      "database (seven tables).")
    body(document,
         "Recommended placement: new Database-design subsection (Component 2.2), next to "
         "src/main/resources/db/schema.sql and the DAO tests. See docs/design/README.md for "
         "the key/cardinality notes.", size=10.5)

    heading(document, "Figure 4 - MVC and layered architecture", page_break=True)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(DESIGN / "figure4_mvc_architecture_p2.png"),
                                    width=Inches(6.2))
    caption(document, "Figure 4: MVC and layered architecture of the Project 2 system with "
                      "the design patterns used in each layer.")
    body(document,
         "Recommended placement: Architecture subsection (Component 2.3.2); the note at the "
         "bottom of the figure maps MVC, DAO, Singleton, Factory, Observer, Strategy and "
         "Unit of Work to the classes.", size=10.5)

    properties = document.core_properties
    properties.title = "Project 2 Design Figures (Report-Ready Draft)"
    properties.author = "Yuheng Hu (26366447)"
    properties.subject = "COMP603/ENSE600 Project 2 - StockFlow Inventory Management System"

    document.save(TARGET)
    reopened = Document(TARGET)
    print("saved", TARGET)
    print("paragraphs:", len(reopened.paragraphs),
          "inline images:", len(reopened.inline_shapes))


if __name__ == "__main__":
    main()
